"""
Document processing module for parsing and chunking documents.
Supports PDF, TXT, and DOCX files.
"""
from pypdf import PdfReader
import docx
import re
from typing import List, Tuple
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


def split_into_sentences(text: str) -> List[str]:
    """
    Split text into sentences using regex patterns.
    Handles common sentence endings: . ? ! and common abbreviations.
    
    Args:
        text: Input text to split
        
    Returns:
        List of sentences
    """
    # Replace common abbreviations with placeholders to avoid false splits
    text = re.sub(r'\bMr\.', 'Mr<PERIOD>', text)
    text = re.sub(r'\bMrs\.', 'Mrs<PERIOD>', text)
    text = re.sub(r'\bMs\.', 'Ms<PERIOD>', text)
    text = re.sub(r'\bDr\.', 'Dr<PERIOD>', text)
    text = re.sub(r'\bProf\.', 'Prof<PERIOD>', text)
    text = re.sub(r'\bSr\.', 'Sr<PERIOD>', text)
    text = re.sub(r'\bJr\.', 'Jr<PERIOD>', text)
    text = re.sub(r'\be\.g\.', 'e<PERIOD>g<PERIOD>', text)
    text = re.sub(r'\bi\.e\.', 'i<PERIOD>e<PERIOD>', text)
    text = re.sub(r'\bInc\.', 'Inc<PERIOD>', text)
    text = re.sub(r'\bLtd\.', 'Ltd<PERIOD>', text)
    text = re.sub(r'\bCo\.', 'Co<PERIOD>', text)
    text = re.sub(r'\bCorp\.', 'Corp<PERIOD>', text)
    text = re.sub(r'\bvs\.', 'vs<PERIOD>', text)
    text = re.sub(r'\betc\.', 'etc<PERIOD>', text)
    
    # Split on sentence endings followed by space and capital letter or newline
    pattern = r'(?<=[.!?])\s+(?=[A-Z])|(?<=[.!?])\n'
    sentences = re.split(pattern, text)
    
    # Restore periods in abbreviations
    sentences = [s.replace('<PERIOD>', '.').strip() for s in sentences if s.strip()]
    
    return sentences


class DocumentChunk:
    """Represents a chunk of text from a document."""
    
    def __init__(
        self,
        text: str,
        chunk_index: int,
        start_char: int,
        end_char: int,
        metadata: dict = None
    ):
        self.text = text
        self.chunk_index = chunk_index
        self.start_char = start_char
        self.end_char = end_char
        self.metadata = metadata or {}


class DocumentProcessor:
    """Process documents: parse and chunk text."""
    
    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 200
    ):
        """
        Initialize document processor.
        
        Args:
            chunk_size: Target size for each chunk (in characters)
            chunk_overlap: Overlap between consecutive chunks (in characters)
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    
    def parse_file(self, file_path: Path, file_type: str) -> str:
        """
        Parse document and extract text.
        
        Args:
            file_path: Path to the document file
            file_type: File type ('pdf', 'txt', 'docx')
            
        Returns:
            Extracted text content
        """
        if file_type == 'pdf':
            return self._parse_pdf(file_path)
        elif file_type == 'txt':
            return self._parse_txt(file_path)
        elif file_type == 'docx':
            return self._parse_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
    
    def _parse_pdf(self, file_path: Path) -> str:
        """Extract text from PDF file."""
        logger.info(f"Parsing PDF: {file_path}")
        text = []
        
        with open(file_path, 'rb') as file:
            pdf_reader = PdfReader(file)
            for page_num, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
        
        return "\n\n".join(text)
    
    def _parse_txt(self, file_path: Path) -> str:
        """Extract text from TXT file."""
        logger.info(f"Parsing TXT: {file_path}")
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    
    def _parse_docx(self, file_path: Path) -> str:
        """Extract text from DOCX file."""
        logger.info(f"Parsing DOCX: {file_path}")
        doc = docx.Document(file_path)
        text = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        
        return "\n\n".join(text)
    
    def chunk_text(self, text: str, metadata: dict = None) -> List[DocumentChunk]:
        """
        Split text into overlapping chunks using sentence-aware boundaries.
        
        Args:
            text: Full document text
            metadata: Optional metadata to attach to each chunk
            
        Returns:
            List of DocumentChunk objects
        """
        if not text or len(text.strip()) == 0:
            return []
        
        # Split text into sentences using our custom sentence splitter
        sentences = split_into_sentences(text)
        
        chunks = []
        current_chunk_sentences = []
        current_chunk_length = 0
        chunk_index = 0
        start_char = 0
        
        for i, sentence in enumerate(sentences):
            sentence_length = len(sentence)
            
            # Check if adding this sentence would exceed chunk_size
            if current_chunk_length + sentence_length > self.chunk_size and current_chunk_sentences:
                # Create chunk from accumulated sentences
                chunk_text = ' '.join(current_chunk_sentences).strip()
                
                if chunk_text:
                    chunk = DocumentChunk(
                        text=chunk_text,
                        chunk_index=chunk_index,
                        start_char=start_char,
                        end_char=start_char + len(chunk_text),
                        metadata=metadata
                    )
                    chunks.append(chunk)
                    chunk_index += 1
                
                # Calculate overlap: keep last few sentences that fit in overlap size
                overlap_sentences = []
                overlap_length = 0
                
                for sent in reversed(current_chunk_sentences):
                    if overlap_length + len(sent) <= self.chunk_overlap:
                        overlap_sentences.insert(0, sent)
                        overlap_length += len(sent) + 1  # +1 for space
                    else:
                        break
                
                # Start new chunk with overlap sentences
                current_chunk_sentences = overlap_sentences
                current_chunk_length = sum(len(s) + 1 for s in overlap_sentences)
                start_char = start_char + len(chunk_text) - overlap_length
            
            # Add current sentence to chunk
            current_chunk_sentences.append(sentence)
            current_chunk_length += sentence_length + 1  # +1 for space between sentences
        
        # Add the last chunk if there are remaining sentences
        if current_chunk_sentences:
            chunk_text = ' '.join(current_chunk_sentences).strip()
            
            if chunk_text:
                chunk = DocumentChunk(
                    text=chunk_text,
                    chunk_index=chunk_index,
                    start_char=start_char,
                    end_char=start_char + len(chunk_text),
                    metadata=metadata
                )
                chunks.append(chunk)
        
        logger.info(f"Created {len(chunks)} sentence-aware chunks from {len(sentences)} sentences (text length: {len(text)})")
        return chunks


# Global document processor instance
_doc_processor: DocumentProcessor = None


def get_document_processor() -> DocumentProcessor:
    """Get or create the global document processor instance."""
    global _doc_processor
    if _doc_processor is None:
        _doc_processor = DocumentProcessor()
    return _doc_processor
